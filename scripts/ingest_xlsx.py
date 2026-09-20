"""ingest_xlsx —— 辅助（不在验证管线内）：回流表 xlsx → INPUT.txt + INPUT.meta.json。
工具集只认 UTF-8 txt；xlsx/csv/json → txt 由本脚本或上游完成，验证工具层做准入。"""
from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date, datetime
from pathlib import Path

import openpyxl

DEFAULT_TEXT_COL = "tts_text"
NUMERIC_COLS = [
    "tts_length", "tts_speed_ratio", "tts_duration", "tts_bit_rate",
    "tts_sample_rate", "tts_s3_key_duration", "tts_user_interrupt_flag",
]


def _cell(v):
    if v is None:
        return None
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    if isinstance(v, (int, float, str, bool)):
        return v
    return str(v)


def _to_num(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    if isinstance(v, str):
        s = v.strip()
        try:
            return int(s)
        except ValueError:
            try:
                return float(s)
            except ValueError:
                return v
    return v


def main(argv=None):
    ap = argparse.ArgumentParser(description="回流表 → INPUT.txt + meta（文本列默认 tts_text）")
    ap.add_argument("--xlsx", required=True, help="回流 xlsx 路径")
    ap.add_argument("--sheet", default=None, help="工作表名（默认第一个；拼错会报错而非静默回退）")
    ap.add_argument("--text-col", default=DEFAULT_TEXT_COL, help="待验证文本列名")
    ap.add_argument("--out-dir", default="input", help="输出目录（默认 input/）")
    ap.add_argument("--lang", default="sw")
    args = ap.parse_args(argv)

    src = Path(args.xlsx)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    txt_path = out_dir / "INPUT.txt"
    meta_path = out_dir / "INPUT.meta.json"

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    # 未指定 --sheet 取第一个；指定了就必须存在——拼错时静默回退到第一个 sheet
    # 会在同名列的情况下读错数据而不报错，故改为明确报错。
    if args.sheet is None:
        ws = wb[wb.sheetnames[0]]
    elif args.sheet in wb.sheetnames:
        ws = wb[args.sheet]
    else:
        raise SystemExit(f"工作表 {args.sheet!r} 不存在。可用的表：{wb.sheetnames}")
    rows = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows)]
    headers = [h for h in header]
    if args.text_col not in headers:
        raise SystemExit(f"文本列 {args.text_col!r} 不在表头：{headers}")
    ti = headers.index(args.text_col)

    lines = []
    n = 0
    with txt_path.open("w", encoding="utf-8", newline="\n") as f:
        for r in rows:
            if r[ti] is None or str(r[ti]).strip() == "":
                continue
            n += 1
            cells = []
            for j, h in enumerate(headers):
                v = _cell(r[j]) if j < len(r) else None
                if h in NUMERIC_COLS:
                    v = _to_num(v)
                cells.append(json.dumps(v, ensure_ascii=False))
            cells.append(json.dumps(str(r[ti]).strip(), ensure_ascii=False))
            f.write("\t".join(cells) + "\n")
    wb.close()

    meta = {
        "source": src.name,
        "sheet": ws.title,
        "lang": args.lang,
        "columns": headers + ["__text__"],
        "text_column": "__text__",
        "numeric_columns": NUMERIC_COLS,
        "rows": n,
        "sha256": hashlib.sha256(src.read_bytes()).hexdigest(),
        "ingested_at": datetime.now().isoformat(timespec="seconds"),
        "note": "INPUT.txt 一行一条；\\t 分列；每格 json 编码；__text__ 为验证文本。",
    }
    meta_path.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"rows": n, "columns": len(headers), "text_col": args.text_col,
                      "txt": str(txt_path), "meta": str(meta_path)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
