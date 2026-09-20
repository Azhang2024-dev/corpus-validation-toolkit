"""隐形信息丢失类缺陷的回归测试（2026-09-18 修复批次）。

统一主题：产物里「有汇总无明细 / 文档承诺未产出 / 静默降级不留痕」——
这类 bug 不会报错，只会让报告上的数字悄悄变错或变少，所以每一条都要有守卫测试。

运行：python -m unittest discover -s tests -v
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
PY = sys.executable

from corpus_kit.charset import check_charset
from corpus_kit.config import Config
from corpus_kit.g2p import G2P
from corpus_kit.io import read_jsonl
from corpus_kit.phonotactics import phonotactics_report

CFG = Config("sw")


class TestGateWarnNotFoldedIntoFail(unittest.TestCase):
    """#1 gate 汇总必须三档分开：原 `n - pass` 把 warn 算成 fail。"""

    def test_charset_warn_and_fail_are_separate_counts(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "run"
            src = Path(td) / "in.txt"
            # 一行纯斯语(pass) + 一行斯语夹外语(warn) + 一行纯外语(fail)
            src.write_text(
                "Habari za asubuhi leo ni siku nzuri sana.\n"
                "Habari za asubuhi, 你好，今天天气很好，我们去公园散步吧?\n"
                "这是一整行中文内容完全不包含目标语言的任何词汇和字符。\n",
                encoding="utf-8")
            r = subprocess.run(
                [PY, str(ROOT / "skills" / "gate_validity" / "run.py"),
                 "--input", str(src), "--out", str(out), "--lang", "sw"],
                capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 0, msg=r.stdout + r.stderr)
            g = json.loads((out / "dims" / "gate.json").read_text(encoding="utf-8-sig"))
            # 三档之和 == 总行数；不得把 warn 折进 fail
            self.assertEqual(g["charset_pass"] + g["charset_warn"] + g["charset_fail"], g["rows"])
            self.assertEqual(g["lang_pass"] + g["lang_warn"] + g["lang_fail"], g["rows"])
            self.assertIn("charset_warn", g)
            self.assertIn("lang_warn", g)
            # 与逐行 records 的真实分布一致
            recs = read_jsonl(out / "records.jsonl")
            from collections import Counter
            cc = Counter(x["gate"]["charset"]["status"] for x in recs)
            self.assertEqual(g["charset_pass"], cc.get("pass", 0))
            self.assertEqual(g["charset_warn"], cc.get("warn", 0))
            self.assertEqual(g["charset_fail"], cc.get("fail", 0))

    def test_report_dims_block_lists_warn_and_fail(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "run"
            r = subprocess.run(
                [PY, str(ROOT / "run_pipeline.py"), "--input", str(ROOT / "tests" / "fixtures" / "sample_sw.txt"),
                 "--lang", "sw", "--out", str(out), "--skip-etymology", "--skip-g2p"],
                capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 0, msg=r.stdout + r.stderr)
            wb = load_workbook(out / "数据集_全字段.xlsx", read_only=True)
            labels = []
            for row in wb["维度分布"].iter_rows(values_only=True):
                if row and row[0]:
                    labels.append(str(row[0]))
            wb.close()
            joined = " | ".join(labels)
            self.assertIn("语言判定 warn 行数", joined)
            self.assertIn("语言判定 fail 行数", joined)
            self.assertIn("字符检查 warn 行数", joined)
            self.assertIn("字符检查 fail 行数", joined)


class TestDegradedLanguagemarked(unittest.TestCase):
    """#2 缺 <lang>.yaml 必须标 degraded，且 resources_ok 不得谎报就绪。"""

    def test_missing_lang_yaml_marks_degraded(self):
        c = Config("zz")                       # 不存在的语种
        self.assertTrue(c.degraded)
        self.assertTrue(c.degraded_reasons)
        self.assertFalse(c.resources_ok, "缺 yaml 时不得判为资源就绪")

    def test_present_lang_not_degraded(self):
        c = Config("sw")
        self.assertFalse(c.degraded)
        self.assertEqual(c.degraded_reasons, [])
        self.assertTrue(c.resources_ok)

    def test_gate_json_carries_degraded_flag(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "run"
            src = Path(td) / "in.txt"
            src.write_text("Habari za asubuhi leo ni siku nzuri sana.\n", encoding="utf-8")
            r = subprocess.run(
                [PY, str(ROOT / "skills" / "gate_validity" / "run.py"),
                 "--input", str(src), "--out", str(out), "--lang", "zz"],
                capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 0, msg=r.stdout + r.stderr)
            g = json.loads((out / "dims" / "gate.json").read_text(encoding="utf-8-sig"))
            self.assertTrue(g["degraded"])
            self.assertTrue(g["degraded_reasons"])


class TestCharsetDiagnosticsComplete(unittest.TestCase):
    """#8 foreign/noise 字符诊断：不得 set() 丢频次、不得 [:20] 截尾。"""

    def test_foreign_chars_keep_multiplicity_and_full_tail(self):
        # 同一越界字符出现多次 + 不同越界字符数 > 20
        text = ("Habari " + "中" * 7 + "".join(chr(0x4E00 + i) for i in range(25)))
        r = check_charset(text, CFG)
        # 频次保留：'中' 应记为 7 次而非去重成 1
        d = dict((c, n) for c, n in r["foreign_chars"])
        self.assertEqual(d.get("中"), 7)
        # 不截断到 20：distinct 至少 26 个
        self.assertEqual(r["foreign_distinct"], len(r["foreign_chars"]))
        self.assertGreaterEqual(r["foreign_distinct"], 26)

    def test_pass_row_has_empty_chars_and_zero_distinct(self):
        r = check_charset("Habari za asubuhi.", CFG)
        self.assertEqual(r["foreign_chars"], [])
        self.assertEqual(r["foreign_distinct"], 0)


class TestReadJsonlWarnsOnBadLines(unittest.TestCase):
    """#7 坏行不得静默丢：要告警且计数可辨。"""

    def test_bad_line_emits_warning(self):
        import io as _io
        import contextlib
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "x.jsonl"
            p.write_text('{"a":1}\nNOT_JSON\n{"a":3}\n', encoding="utf-8")
            buf = _io.StringIO()
            with contextlib.redirect_stderr(buf):
                got = read_jsonl(p, name="x.jsonl")
            self.assertEqual(len(got), 2)
            self.assertIn("无法解析", buf.getvalue())
            self.assertIn("x.jsonl", buf.getvalue())

    def test_clean_file_no_warning(self):
        import io as _io
        import contextlib
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "y.jsonl"
            p.write_text('{"a":1}\n{"a":2}\n', encoding="utf-8")
            buf = _io.StringIO()
            with contextlib.redirect_stderr(buf):
                got = read_jsonl(p)
            self.assertEqual(len(got), 2)
            self.assertEqual(buf.getvalue(), "")


class TestPhonotacticsNotTruncated(unittest.TestCase):
    """#4 out_of_space 必须全量落盘（原 [:120] 静默丢弃尾部）。"""

    def test_all_out_of_space_listed(self):
        from collections import Counter
        # 造 130 个越界组合（远超原 120 上限）
        pairs = Counter()
        for i in range(130):
            pairs[("ʘ", f"q{i}")] = i + 1
        rep = phonotactics_report(pairs, CFG)
        self.assertEqual(rep["n_out_of_space"], 130)
        self.assertEqual(len(rep["out_of_space"]), 130, "越界组合应全量列出，不得截断")


class TestG2pMissAttribution(unittest.TestCase):
    """#9 未覆盖要能归因，不能只有个数。"""

    def test_miss_breakdown_distinguishes_causes(self):
        g = G2P(CFG)
        # "123"/"12ab" 非词形（含数字，WORD_RE 不匹配）；"zzzzqqqq" 是词形但不在词典
        words = ["habari", "123", "12ab", "zzzzqqqq"]
        b = g.miss_breakdown(words)
        self.assertEqual(set(b), {"not_a_word_shape", "no_fallback_configured",
                                  "epitran_failed", "epitran_error"})
        self.assertGreaterEqual(b["not_a_word_shape"], 2)

    def test_successful_epitran_not_counted_as_failed(self):
        """能转写出来的词是已覆盖，绝不能记进 epitran_failed。"""
        g = G2P(CFG)
        b = g.miss_breakdown(["habari", "sasa", "zzzzqqqq"])
        # 只要 epitran 可用，"zzzzqqqq" 会被转写（非 None），epitran_failed 必须是 0
        if g.use_epitran:
            self.assertEqual(b["epitran_failed"], 0,
                             "epitran 转写成功的词被误记为故障")

    def test_fallback_none_disables_epitran(self):
        from corpus_kit.config import Config as _C
        cfg = _C("sw")
        cfg.g2p_fallback = "none"
        g = G2P(cfg)
        self.assertFalse(g.use_epitran)
        # 纯字母词、不在词典 → 无兜底器即未覆盖，且归因为 no_fallback（不是 epitran_failed）
        b = g.miss_breakdown(["zzzzqqqq"])
        self.assertGreaterEqual(b["no_fallback_configured"], 1)
        self.assertEqual(b["epitran_failed"], 0)


class TestManifestRecordsRunArgs(unittest.TestCase):
    """#5 manifest 必须记本次实际参数，否则 --limit 批次与全量无法区分。"""

    def test_limit_recorded_in_manifest(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "run"
            r = subprocess.run(
                [PY, str(ROOT / "run_pipeline.py"), "--input", str(ROOT / "tests" / "fixtures" / "sample_sw.txt"),
                 "--lang", "sw", "--out", str(out), "--skip-etymology", "--skip-g2p", "--limit", "2"],
                capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 0, msg=r.stdout + r.stderr)
            m = json.loads((out / "manifest.json").read_text(encoding="utf-8-sig"))
            self.assertIn("run_args", m["params"])
            self.assertEqual(m["params"]["run_args"]["limit"], 2)
            self.assertTrue(m["params"]["run_args"]["skip_g2p"])


class TestKeepQuarantinedCountsConsistent(unittest.TestCase):
    """#6 --keep-quarantined 时 manifest 计数须与 INPUT.txt 实际行数一致。"""

    def test_counts_match_input_lines(self):
        with tempfile.TemporaryDirectory() as td:
            d = Path(td)
            src = d / "in.txt"
            src.write_text("\n".join([
                "Habari za asubuhi leo ni siku nzuri sana.",
                "Habari za asubuhi leo ni siku nzuri sana.",   # 重复
                "!!! ??? ...",                                  # 硬剔除
                "Mambo ya leo ni mazuri kabisa.",               # 软/正常
            ]) + "\n", encoding="utf-8")
            out = d / "o"
            r = subprocess.run(
                [PY, str(ROOT / "scripts" / "prepare_input.py"),
                 "--input", str(src), "--lang", "sw", "--out-dir", str(out), "--keep-quarantined"],
                capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 0, msg=r.stdout + r.stderr)
            n_lines = len((out / "INPUT.txt").read_text(encoding="utf-8").splitlines())
            meta = json.loads((out / "INPUT.meta.json").read_text(encoding="utf-8-sig"))
            cm = json.loads((out / "clean_manifest.json").read_text(encoding="utf-8-sig"))
            self.assertEqual(meta["rows"], n_lines)
            self.assertEqual(meta["clean"]["n_kept"], n_lines)
            self.assertEqual(cm["counts"]["n_kept"], n_lines,
                             "clean_manifest.counts.n_kept 必须等于 INPUT.txt 实际行数")


class TestVocabTopTruncationSelfEvident(unittest.TestCase):
    """#3 Top100 截断要自证，不能让读者以为 top 就是全部型符。"""

    def test_top_flags_present(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td)
            sents = [{"sent_id": 1, "text": "habari za asubuhi"}]
            (out / "records.jsonl").write_text(
                json.dumps({"row_id": 1, "text": "habari za asubuhi", "fields": []}) + "\n", encoding="utf-8")
            (out / "sentences.jsonl").write_text(
                json.dumps(sents[0], ensure_ascii=False) + "\n", encoding="utf-8")
            r = subprocess.run(
                [PY, str(ROOT / "skills" / "profile_vocab" / "run.py"),
                 "--records", str(out / "records.jsonl"), "--sentences", str(out / "sentences.jsonl"),
                 "--lang", "sw", "--out", str(out)],
                capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 0, msg=r.stdout + r.stderr)
            d = json.loads((out / "dims" / "vocab.json").read_text(encoding="utf-8-sig"))
            self.assertIn("top_n", d)
            self.assertIn("top_truncated", d)
            self.assertIn("top_full_source", d)
            # 只有 3 个型符 → 未截断
            self.assertFalse(d["top_truncated"])
            self.assertEqual(d["top_n"], d["types"])


if __name__ == "__main__":
    raise SystemExit(unittest.main())