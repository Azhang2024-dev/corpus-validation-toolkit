"""V4.3/V4.4/V2 回归：音素组合空间与缺失组合 / G2P 一致性抽验 / 句式偏差告警。

对应 2026-09-16 补做的三项（框架 14 子项中判定为「有下游用途」的实缺）：
  (1) V4.3 trigram 未做，但补齐了「理论组合空间 + 缺失组合清单」——选句补齐的输入
  (2) V4.4 G2P 一致性抽验——四档口径（exact/归一/折合表示层/判定）
  (3) V2 句式偏差告警——与锁定目标配比的差距 + 告警

运行：python -m unittest discover -s tests -v
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
PY = sys.executable

from corpus_kit.config import Config
from corpus_kit.g2p import G2P, load_ipa_dict
from corpus_kit.g2p_audit import audit, fold_sequence
from corpus_kit.phonemes import phoneme_dist
from corpus_kit.phonotactics import (
    combo_space,
    normalize_phoneme,
    normalize_sequence,
    phonotactics_report,
)

CFG = Config("sw")


class TestNormalizePhonemes(unittest.TestCase):
    """异写归并 + 前鼻化拆分：统计前必须归一，否则音素清单/组合空间被污染。"""

    def test_equivalence_merge(self):
        self.assertEqual(normalize_phoneme("\u0261", CFG), "g")   # ɡ → g
        self.assertEqual(normalize_phoneme("r", CFG), "\u027e")   # r → ɾ

    def test_prenasal_decompose(self):
        # 上标前鼻化单音位拆回 N+C
        self.assertEqual(normalize_sequence("\u1d50\u0253", CFG), ["m", "\u0253"])   # ᵐɓ
        self.assertEqual(normalize_sequence("\u207f\u0257", CFG), ["n", "\u0257"])   # ⁿɗ
        self.assertEqual(normalize_sequence("\u207fz", CFG), ["n", "z"])             # ⁿz
        self.assertEqual(normalize_sequence("kwa", CFG), ["k", "w", "a"])

    def test_merge_removes_variant_glyph(self):
        """g/ɡ 归并后不应同时出现两个变体；前鼻化拆分后不应残留上标形式。"""
        from corpus_kit.phonemes import phoneme_dist
        vals = ["ha\u0261a", "\u0261a\u027ei"]   # 含 ɡ
        norm = phoneme_dist(vals, normalize=lambda s: normalize_sequence(s, CFG))["single"]
        self.assertIn("g", norm)
        self.assertNotIn("\u0261", norm)          # ɡ 已被归并
        vals2 = ["\u1d50\u0253uzi"]                # ᵐɓ
        norm2 = phoneme_dist(vals2, normalize=lambda s: normalize_sequence(s, CFG))["single"]
        self.assertIn("m", norm2)
        self.assertIn("\u0253", norm2)
        self.assertNotIn("\u1d50\u0253", norm2)   # 上标形式已拆分


class TestComboSpace(unittest.TestCase):
    """理论组合空间 + 缺失/越界。"""

    def test_space_contains_core_types(self):
        sp = combo_space(CFG)
        self.assertIn(("k", "a"), sp)      # CV
        self.assertIn(("a", "k"), sp)      # VC
        self.assertIn(("a", "i"), sp)      # VV
        self.assertIn(("m", "b"), sp)      # 鼻音丛（同化：m 接唇音）
        self.assertIn(("k", "w"), sp)      # C+Glide
        self.assertGreater(len(sp), 100)

    def test_missing_is_space_minus_observed(self):
        obs = Counter({("k", "a"): 3, ("a", "i"): 2})
        rep = phonotactics_report(obs, CFG)
        self.assertEqual(rep["n_observed"], 2)
        self.assertEqual(rep["n_missing"], rep["n_space"] - 2)
        combos = {m["combo"] for m in rep["missing"]}
        self.assertNotIn("ka", combos)   # 已观测，不该出现在 missing

    def test_out_of_space_detected(self):
        # 外来丛 st 不在 native 空间
        rep = phonotactics_report(Counter({("s", "t"): 5}), CFG)
        self.assertEqual(rep["n_out_of_space"], 1)
        self.assertEqual(rep["out_of_space"][0]["combo"], "st")
        self.assertEqual(rep["out_of_space"][0]["count"], 5)

    def test_word_internal_only_no_cross_word(self):
        """组合必须来自词内序列；跨词拼接会造伪组合（如 ab）。"""
        # habari → h a ɓ a ɾ i，不应产生 a+b
        seq = normalize_sequence("ha\u0253a\u027ei", CFG)
        bigrams = set(zip(seq, seq[1:]))
        self.assertIn(("a", "\u0253"), bigrams)
        self.assertNotIn(("a", "b"), bigrams)


class TestG2PAudit(unittest.TestCase):
    """四档口径：exact / 归一 / 折合表示层 / 判定。"""

    def test_fold_sequence_merges_conventions(self):
        # n g 与 ŋ g 折合后应视为同一标记
        folds = (CFG.get("g2p", {}) or {}).get("audit_fold")
        a = fold_sequence(["n", "g", "a"], folds)
        b = fold_sequence(["\u014b", "g", "a"], folds)
        self.assertEqual(a, b)

    def test_audit_real_data_folds_to_pass(self):
        """Honesty after narrowing audit_fold: dual-gate, no fold-only greenwash.

        Notation-only fold (n g ↔ ŋ g) may still lift rate_folded above
        rate_phoneme, but phonemic contrasts (implosives etc.) must remain
        visible in rate_phoneme. Overall status is the worse of the two gates.
        """
        g = G2P(CFG)
        d = load_ipa_dict(CFG.ipa_dict_path)
        r = audit(g, d, CFG, n=100, seed=0)
        self.assertEqual(r["n_sample"], 100)
        # Dual-gate fields always present
        for key in ("status_folded", "status_phoneme", "status_overall", "status"):
            self.assertIn(key, r)
            self.assertIn(r[key], {"pass", "warn"})
        # status mirrors overall (cannot greenwash on fold alone)
        self.assertEqual(r["status"], r["status_overall"])
        expected_overall = (
            "pass"
            if r["status_folded"] == "pass" and r["status_phoneme"] == "pass"
            else "warn"
        )
        self.assertEqual(r["status_overall"], expected_overall)
        # Fold may still help a little (ŋg notation), but must not erase phonemic gaps
        self.assertGreaterEqual(r["rate_folded"], r["rate_phoneme"])
        # With phonemic pairs removed from audit_fold, a fold-only 0.95 pass
        # is no longer required; phoneme gate is the honesty check.
        if r["rate_folded"] >= r["threshold"] and r["rate_phoneme"] < r.get("threshold_phoneme", r["threshold"]):
            self.assertEqual(r["status_folded"], "pass")
            self.assertEqual(r["status_phoneme"], "warn")
            self.assertEqual(r["status_overall"], "warn")
        # 抽验产物结构
        self.assertIn("mismatches", r)
        for m in r["mismatches"]:
            self.assertIn("word", m)
            self.assertIn("dict_ipa", m)

    def test_audit_fold_excludes_phonemic_contrasts(self):
        """sw.yaml audit_fold must not equate real phonemic contrasts."""
        folds = (CFG.get("g2p", {}) or {}).get("audit_fold") or []

        def _norm_side(side):
            # YAML may load "n g" as a string or as ["n", "g"]
            if isinstance(side, str):
                tokens = side.split()
            else:
                tokens = []
                for part in side:
                    tokens.extend(str(part).split())
            return " ".join(tokens)

        normalized = {tuple(sorted((_norm_side(a), _norm_side(b)))) for a, b in folds}
        forbidden = {
            tuple(sorted(("m b", "m ɓ"))),
            tuple(sorted(("n d", "n ɗ"))),
            tuple(sorted(("n ɟ", "ɲ dʒ"))),
            tuple(sorted(("t h", "θ"))),
        }
        self.assertTrue(normalized.isdisjoint(forbidden),
                        f"phonemic folds still present: {normalized & forbidden}")
        # notation-only velar fold may remain
        self.assertIn(tuple(sorted(("n g", "ŋ g"))), normalized)

    def test_audit_threshold_from_config(self):
        g = G2P(CFG)
        d = load_ipa_dict(CFG.ipa_dict_path)
        r = audit(g, d, CFG, n=50, seed=1)
        self.assertEqual(r["threshold"], 0.95)


class TestSentenceWarning(unittest.TestCase):
    """V2 目标配比差距 + 告警。"""

    def _dims(self, texts, out):
        from corpus_kit.io import write_jsonl
        from corpus_kit.tokenize import word_tokens
        sents = [{"sent_id": i + 1, "row_id": 1, "text": t, "tokens": word_tokens(t, r"[A-Za-z'\-]+")}
                 for i, t in enumerate(texts)]
        sp = Path(out) / "sentences.jsonl"
        write_jsonl(sp, sents)
        r = subprocess.run(
            [PY, str(ROOT / "skills" / "profile_sentence" / "run.py"),
             "--sentences", str(sp), "--lang", "sw", "--out", str(out)],
            capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(r.returncode, 0, msg=r.stdout + r.stderr)
        return json.loads((Path(out) / "dims" / "sentence.json").read_text(encoding="utf-8-sig"))

    def test_warn_when_below_target(self):
        with tempfile.TemporaryDirectory() as td:
            # 全陈述句 → 疑问/感叹/短语都偏低；陈述 100% 超上限 → 偏高
            d = self._dims(["Hii ni sentensi ya kawaida."] * 20, td)
        self.assertEqual(d["status"], "warn")
        self.assertGreaterEqual(d["n_warn"], 1)
        by = {x["句式"]: x for x in d["distribution"]}
        self.assertEqual(by["陈述"]["告警"], "偏高")   # 100% > 80% 上限
        self.assertEqual(by["疑问"]["告警"], "偏低")   # 0% < 10% 下限

    def test_pass_when_in_range(self):
        with tempfile.TemporaryDirectory() as td:
            # 陈述 75%、疑问 15%、感叹 10%、短语 0 → 短语偏低但其余合格
            texts = ["Hii ni sentensi ya kawaida nzuri."] * 15 + ["Je, uko tayari?"] * 3 + ["Nzuri sana!"] * 2
            d = self._dims(texts, td)
        by = {x["句式"]: x for x in d["distribution"]}
        self.assertEqual(by["陈述"]["告警"], "")
        self.assertEqual(by["疑问"]["告警"], "")
        self.assertEqual(by["感叹"]["告警"], "")


class TestReportSheets(unittest.TestCase):
    """报告含新 sheet 与告警列。"""

    def test_end_to_end_has_new_sheets(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "run"
            r = subprocess.run(
                [PY, str(ROOT / "run_pipeline.py"),
                 "--lang", "sw",
                 "--input", str(ROOT / "tests" / "fixtures" / "sample_sw.txt"),
                 "--out", str(out), "--skip-etymology", "--limit", "8"],
                capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 0, msg=r.stdout + r.stderr)
            self.assertTrue((out / "g2p_audit.csv").exists())
            wb = load_workbook(out / "数据集_全字段.xlsx", read_only=True)
            for name in ("音素组合", "缺失组合", "G2P抽验"):
                self.assertIn(name, wb.sheetnames, f"缺 sheet {name}")
            wb.close()
            # manifest 记录三项状态
            m = json.loads((out / "manifest.json").read_text(encoding="utf-8-sig"))
            self.assertIn("missing_combos", m["input"])
            self.assertIn("g2p_audit_status", m["input"])
            self.assertIn("sentence_status", m["input"])

    def test_csv_has_single_header_row(self):
        """回归：全字段 CSV 表头只能有一行（原实现把 worksheet 第 1 行表头又写了一遍）。"""
        import csv as _csv
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "run"
            r = subprocess.run(
                [PY, str(ROOT / "run_pipeline.py"),
                 "--lang", "sw",
                 "--input", str(ROOT / "tests" / "fixtures" / "sample_sw.txt"),
                 "--out", str(out), "--skip-etymology", "--skip-g2p", "--limit", "6"],
                capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 0, msg=r.stdout + r.stderr)
            with (out / "数据集_全字段.csv").open(encoding="utf-8-sig", newline="") as f:
                rows = list(_csv.reader(f))
        self.assertEqual(rows[0][0], "sent_id")            # 第 1 行是表头
        self.assertNotEqual(rows[1][0], "sent_id")         # 第 2 行必须是数据
        self.assertEqual(rows[1][0], "1")                  # sent_id 从 1 起

    def test_missing_combo_header_style_on_header_row(self):
        """回归：「缺失组合」sheet 的表头样式须加在表头行，不是第一条数据行。"""
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "run"
            r = subprocess.run(
                [PY, str(ROOT / "run_pipeline.py"),
                 "--lang", "sw",
                 "--input", str(ROOT / "tests" / "fixtures" / "sample_sw.txt"),
                 "--out", str(out), "--skip-etymology", "--limit", "8"],
                capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(r.returncode, 0, msg=r.stdout + r.stderr)
            wb = load_workbook(out / "数据集_全字段.xlsx")
            ws = wb["缺失组合"]
            # 找到表头文字所在行
            hdr_row = next(r_ for r_ in range(1, ws.max_row + 1)
                           if ws.cell(r_, 1).value == "缺失组合（选句补齐候选）")
            # 表头单元格必须有深色底（HEAD_FILL），且下一行（数据）不应带该底色
            self.assertIsNotNone(ws.cell(hdr_row, 1).fill.fgColor.rgb)
            self.assertNotEqual(str(ws.cell(hdr_row, 1).fill.fgColor.rgb), "00000000")
            self.assertNotEqual(str(ws.cell(hdr_row + 1, 1).fill.fgColor.rgb), "001F4E79")
            wb.close()

    def test_sheet2_median_is_standard(self):
        """回归：Sheet2「中位句长」须为标准中位数（偶数上中位是旧口径）。"""
        from corpus_kit.report import build_report
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "run"
            out.mkdir()
            # 单行 4 句、字符数 1/2/3/4 → 标准中位 2.5；旧实现会得 3
            sents = [{"sent_id": i + 1, "row_id": 1, "text": "x" * (i + 1), "n_chars": i + 1,
                      "n_words": 1, "n_syllables": 0, "tokens": ["x"]} for i in range(4)]
            records = [{"row_id": 1, "text": "xxxx", "fields": [], "gate": {}}]
            build_report(out, {"columns": [], "text_column": None}, sents, records, {}, {})
            wb = load_workbook(out / "数据集_全字段.xlsx", read_only=True)
            ws = wb["行级汇总"]
            headers = [c.value for c in next(ws.iter_rows(max_row=1))]
            med = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))][headers.index("中位句长")]
            wb.close()
        self.assertEqual(med, 2.5)


if __name__ == "__main__":
    unittest.main()