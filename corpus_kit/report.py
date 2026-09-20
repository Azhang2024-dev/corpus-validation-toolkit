"""report —— 落盘：数据集_全字段.xlsx（原字段保留+维度列+多 sheet）+ csv + 词级明细 + manifest.json。

词级明细（2026-09-16 新增）
-------------------------
计算层（`token_ipa.json` 词→IPA、`etymology.jsonl` 词→词源）本就齐备，此前报告层没有 join，
导致「每个词的 IPA / 词源」在交付物里不可见。本模块补上 join：
  · Sheet1 句子级明细 追加 `词级明细` 列：`词[IPA|词源] 词[IPA|词源] ...`（逐词可见）
  · Sheet5 `词级明细`：去重词表，逐词一行（word / ipa / 音素 / 词源 / 频次）
  · `数据集_词级明细.csv` + `词级明细.jsonl`：同一份词级数据的程序友好形态
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import statistics
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import __version__
from .io import ensure_dir, read_jsonl, write_jsonl
from .phonemes import split_phonemes

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True)
WARN_FONT = Font(color="9C6500")
FAIL_FONT = Font(color="C00000")

# 新增维度列说明（字段说明子表）
NEW_COLUMNS = [
    ("sent_id", "句级唯一 ID（1..N，贯穿全链路）", "行"),
    ("row_id", "原 INPUT.txt 行号（1..N），回溯原回流表", "行"),
    ("gate_lang", "目标语言判定：pass / warn / fail", "文本"),
    ("gate_charset", "字符检查：pass / warn / fail（越界脚本比例分桶）", "文本"),
    ("gate_reason", "gate 判定原因摘要（虚词命中率 / 越界比例等）", "文本"),
    ("n_chars", "单句字符数 len()", "整数"),
    ("n_words", "单句词数（空白切词）", "整数"),
    ("n_syllables", "单句音节数（拼写元音 aeiou 计数，与 IPA 100% 一致）", "整数"),
    ("sentence_type", "句式：疑问 / 感叹 / 陈述 / 短语", "文本"),
    ("nsw_types", "命中 NSW 类别（逗号分隔；无则空）", "文本"),
    ("nsw_hits", "命中 NSW 的具体串：`类别:串1,串2; 类别2:…`（无则空）", "文本"),
    ("n_phonemes", "句内可音素化词的音素总数（查表+epitran 兜底后）", "整数"),
    ("phoneme_types", "句内不同音素数", "整数"),
    ("词级明细", "逐词展开：`词[IPA|词源]`，空格分隔；`-` 表示该词 IPA/词源缺失", "文本"),
    ("etymology_status", "词源打标：tagged / partial / untagged / etym_fail", "文本"),
]

# 词级明细 sheet / csv / jsonl 的列
WORD_COLUMNS = [
    ("word", "词（小写，词法切词口径）", "文本"),
    ("ipa", "该词 IPA（查表→epitran→补丁；空=未覆盖）", "文本"),
    ("phonemes", "IPA 切分后的音素序列（空格分隔）", "文本"),
    ("n_phonemes", "音素数", "整数"),
    ("origin", "词源：bantu_native / arabic_loan / english_loan / other_loan / uncertain", "文本"),
    ("count", "该词在全批语料中的 token 频次", "整数"),
]

# 回流原表头说明（已知 23 列；未知列自动兜底『原表字段』）
ORIGINAL_COLUMNS = {
    "text_query": "用户查询原文（JSON 数组字符串）",
    "tts_model_version": "模型版本",
    "tts_voice_type": "音色类型（如 sw-KE-ZuriNeural）",
    "tts_model_name": "语音合成模型/厂商",
    "tts_s3_key": "音频 S3 key",
    "p_date": "分区日期",
    "tts_bit_rate": "比特率",
    "tts_length": "原始字段-文本长度",
    "tts_sample_rate": "采样率",
    "tts_encoding": "编码格式（opus 等）",
    "tts_s3_key_duration": "S3 音频时长（ms）",
    "tts_answer": "助手回复文本（JSON 数组字符串）",
    "tts_speed_ratio": "语速比例",
    "tts_duration": "音频时长（ms）",
    "tts_vendor": "供应商",
    "tts_synthesis_method": "合成方式（online 等）",
    "tts_text": "实际合成文本（TTS 录音用，验证主体）",
    "tts_language": "语言码（sw）",
    "dm_asr_info_is_in_tts": "ASR 信息是否在 TTS 中",
    "tts_locale_language": "语言区域（sw_KE/sw_TZ）",
    "multimodal_query": "多模态查询 JSON",
    "request_id": "请求唯一 ID",
    "tts_user_interrupt_flag": "用户打断标记",
}


def _hash_file(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def _col_widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _header(ws, headers: list[str]):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT


def _tag(ws, r: int, col: int, status: str):
    cell = ws.cell(row=r, column=col)
    if status == "fail":
        cell.font = FAIL_FONT
    elif status == "warn":
        cell.font = WARN_FONT


def _num(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        try:
            return float(v)
        except (TypeError, ValueError):
            return v


def load_word_ipa(outdir: Path) -> dict:
    """token_ipa.json：{词: IPA|None}。文件缺失时返回空表（报告仍可出，只是词级 IPA 为空）。"""
    p = outdir / "token_ipa.json"
    if not p.exists():
        return {}
    try:
        d = json.loads(p.read_text(encoding="utf-8-sig"))
        return d if isinstance(d, dict) else {}
    except Exception:
        return {}


def load_etymology(outdir: Path) -> tuple[dict, dict]:
    """etymology.jsonl → (词→词源, (sent_id,word_id)→词源)。

    兼容两种口径：新词级行 `{word,origin}`；旧句级行 `{sent_id,word_id,word,origin}`。
    新行优先（全局口径更一致）。"""
    by_word: dict = {}
    by_sent_word: dict = {}
    p = outdir / "etymology.jsonl"
    if not p.exists():
        return by_word, by_sent_word
    for line in p.read_text(encoding="utf-8-sig").splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            rec = json.loads(line)
        except json.JSONDecodeError:
            continue
        w = rec.get("word")
        o = rec.get("origin")
        if not w or not o:
            continue
        if "sent_id" in rec and rec.get("word_id") is not None:
            by_sent_word[(rec["sent_id"], rec["word_id"])] = o
        if "sent_id" not in rec:
            by_word[w] = o
    return by_word, by_sent_word


def _phonemes_str(ipa) -> str:
    if not ipa:
        return ""
    return " ".join(split_phonemes(ipa))


def fmt_nsw_hits(hits) -> str:
    """nsw_hits dict → `类别:串1,串2; 类别2:串3`（无则空串）。"""
    if not hits:
        return ""
    return "; ".join(f"{name}:{','.join(vals)}" for name, vals in hits.items())


def fmt_word_detail(tokens, ipa_map, ety_by_word, ety_by_sent_word, sent_id) -> str:
    """把一句的 tokens 展开为 `词[IPA|词源]` 串（逐词可见）。"""
    parts = []
    for i, t in enumerate(tokens or []):
        ipa = ipa_map.get(t) or "-"
        origin = ety_by_word.get(t) or ety_by_sent_word.get((sent_id, i)) or "-"
        parts.append(f"{t}[{ipa}|{origin}]")
    return " ".join(parts)


def build_word_rows(outdir: Path, sentences: list[dict]) -> list[dict]:
    """去重词表 → 逐词一行（word/ipa/phonemes/n_phonemes/origin/count）。"""
    ipa_map = load_word_ipa(outdir)
    ety_by_word, _ = load_etymology(outdir)
    cnt: dict = {}
    for s in sentences:
        for t in s.get("tokens") or []:
            if t:
                cnt[t] = cnt.get(t, 0) + 1
    rows = []
    for w, c in sorted(cnt.items(), key=lambda kv: (-kv[1], kv[0])):
        ipa = ipa_map.get(w)
        rows.append({
            "word": w,
            "ipa": ipa or "",
            "phonemes": _phonemes_str(ipa),
            "n_phonemes": len(split_phonemes(ipa)) if ipa else 0,
            "origin": ety_by_word.get(w, ""),
            "count": c,
        })
    return rows


def _desc_overrides(cfg) -> dict:
    """字段说明文案：sw 保持历史文案（基线冻结），其他语言按配置生成。

    这些文案描述的是「原表字段/新增字段」的含义，与数字无关——不改 sw 的历史输出，
    非 sw 语言不再挂着斯语示例。
    """
    if cfg is None:
        return {}
    lang = str(getattr(cfg, "lang", "") or "")
    if lang == "sw":
        return {}
    method = str(getattr(cfg, "syllable_method", "vowel_groups"))
    syll = {
        "vowel_groups": "单句音节数（拼写元音计数；该语言正字法为音节性拼写）",
        "dict": "单句音节数（词典真值；未覆盖词回退元音组近似，覆盖率见维度分布）",
        "from_ipa": "单句音节数（由 IPA 音素数推算）",
        "na": "单句音节数（N/A：该语言音节口径未定稿，本报告不出数）",
    }.get(method, "单句音节数")
    origins = " / ".join(getattr(cfg, "ety_valid_origins", None) or ["uncertain"])
    return {
        "n_syllables": syll,
        "origin": f"词源：{origins}",
        "tts_language": f"语言码（{lang}）",
        "tts_locale_language": f"语言区域（如 {lang}_<区域>）",
        "tts_voice_type": f"音色类型（如 {lang}-<区域>-<音色名>）",
    }


def build_capability_rows(cfg, gate: dict, length: dict, ph: dict, run_params: dict) -> list:
    """「能力与口径」块：本次计算用到的资源、检测结论、未算维度与原因。

    诚实原则的落点：报告读者第一眼就能看到「这份数字建立在什么之上、哪里是缺的」，
    不必读完每个 sheet 才知道哪些维度不可用。
    """
    rows: list = [["项", "值"]]
    lang = run_params.get("lang") or (getattr(cfg, "lang", "") if cfg else "")
    if cfg is not None:
        rows.append(["目标语言", f"{cfg.label}（{cfg.code}）"])
        rows.append(["脚本家族", f"{cfg.script}（{cfg.script_label}）"])
    else:
        rows.append(["目标语言", str(lang)])
    lc = (gate or {}).get("lang_check") or {}
    if lc:
        st = lc.get("status")
        if st == "auto":
            rows.append(["语言识别（auto）",
                         f"→ {lc.get('picked')}（依据：{lc.get('method')}；"
                         f"检测脚本 {lc.get('detected_script')}，LID {lc.get('detected_lang')}）"])
        elif st == "conflict":
            rows.append(["⚠ 指令与数据不符", "；".join(lc.get("notes") or [])])
        elif st == "unverifiable":
            rows.append(["语言核对", "；".join(lc.get("notes") or []) or "无法核对"])
        elif st == "ok":
            note = "；".join(lc.get("notes") or [])
            rows.append(["语言核对", "指令与数据一致" + (f"（{note}）" if note else "")])
    if cfg is not None:
        dp = getattr(cfg, "ipa_dict_path", None)
        rows.append(["发音词典", (f"{dp.name}（已加载）" if dp and Path(dp).exists()
                                   else "无——音素维度只能靠 G2P 兜底或不可用（见缺口）")])
        rows.append(["G2P 兜底", (f"epitran {cfg.epitran_lang}" if cfg.epitran_lang else "未配置")])
        syl_method = str(getattr(cfg, "syllable_method", "vowel_groups"))
        cov = (length or {}).get("syll_coverage") or {}
        syl_note = {
            "vowel_groups": "按拼写元音组近似",
            "dict": (f"词典真值；未覆盖词回退近似（命中率 "
                     f"{cov.get('dict_hit_ratio') if cov.get('dict_hit_ratio') is not None else '—'}）"),
            "from_ipa": "由 IPA 音素数推算",
            "na": "N/A：口径未定稿，本语言不出音节数",
        }.get(syl_method, syl_method)
        rows.append(["音节口径", syl_note])
        ch_mode = (length or {}).get("char_unit")
        if ch_mode:
            rows.append(["字符口径", "字素簇（\\X）" if ch_mode == "grapheme" else "码点（len()，sw 历史口径）"])
        fw = (gate or {}).get("function_words_available")
        rows.append(["虚词表", "已配置" if fw else "未配置——语言判定仅依赖 langdetect（短文本不稳）"])
    au = (ph or {}).get("g2p_audit") or {}
    if au:
        rows.append(["G2P 抽验", f"{au.get('status')}"
                     + (f"（折合一致率 {au.get('rate_folded')}）" if au.get("rate_folded") is not None else "")
                     + (f"；{au.get('note')}" if au.get("note") else "")])
    ptr = (ph or {}).get("phonotactics") or {}
    if ptr and ptr.get("n_space") == 0:
        rows.append(["音素组合空间", "未配置（该语言无音位学定义）——组合/缺失维度不可用"])
    gaps = list((gate or {}).get("capability_gaps") or [])
    if not gaps and cfg is not None:
        gaps = list(getattr(cfg, "capability_gaps", []) or [])
    if gaps:
        for g in gaps:
            rows.append(["⚠ 能力缺口", g])
    else:
        rows.append(["能力缺口", "无（本语言资源齐备）"])
    return rows


def build_report(outdir: Path, meta: dict, sentences: list[dict], records: list[dict],
                 dims: dict, run_params: dict, cfg=None) -> dict:
    outdir = ensure_dir(outdir)
    rec_by_id = {r["row_id"]: r for r in records}
    src_cols = [c for c in (meta.get("columns") or []) if c != meta.get("text_column")]
    numeric_cols = set(meta.get("numeric_columns") or [])

    ipa_map = load_word_ipa(outdir)
    ety_by_word, ety_by_sent_word = load_etymology(outdir)
    word_rows = build_word_rows(outdir, sentences)

    wb = Workbook()

    # ---------- Sheet1 句子级明细 ----------
    ws = wb.active
    ws.title = "句子级明细"
    headers = ["sent_id", "row_id"] + src_cols + ["文本"] + [
        "gate_lang", "gate_charset", "gate_reason",
        "n_chars", "n_words", "n_syllables", "sentence_type", "nsw_types", "nsw_hits",
        "n_phonemes", "phoneme_types", "词级明细", "etymology_status",
    ]
    _header(ws, headers)
    gate_lang_col = headers.index("gate_lang") + 1
    gate_charset_col = headers.index("gate_charset") + 1
    r = 1
    for s in sentences:
        rec = rec_by_id.get(s.get("row_id"), {})
        gate = rec.get("gate", {}) or {}
        gl = gate.get("lang") or {}
        gc = gate.get("charset") or {}
        row = [
            s.get("sent_id"), s.get("row_id"),
        ]
        src_pos = {c: i for i, c in enumerate(src_cols)}
        for c in src_cols:
            fields = rec.get("fields") or []
            idx = src_pos[c]
            v = fields[idx] if idx < len(fields) else ""
            if c in numeric_cols:
                v = _num(v)
            row.append(v)
        row.append(s.get("text", ""))
        row += [
            gl.get("status", ""), gc.get("status", ""),
            "; ".join(x for x in [gl.get("reason"), gc.get("reason")] if x),
            s.get("n_chars"), s.get("n_words"), s.get("n_syllables"),
            s.get("sentence_type", ""),
            ",".join(s.get("nsw_types") or []),
            fmt_nsw_hits(s.get("nsw_hits")),
            s.get("n_phonemes", 0), s.get("phoneme_types", 0),
            fmt_word_detail(s.get("tokens"), ipa_map, ety_by_word, ety_by_sent_word, s.get("sent_id")),
            s.get("etymology_status", "untagged"),
        ]
        ws.append(row)
        r += 1
        _tag(ws, r, gate_lang_col, gl.get("status", ""))
        _tag(ws, r, gate_charset_col, gc.get("status", ""))
    _col_widths(ws, {"A": 10, "B": 10})
    ws.freeze_panes = "C2"

    # ---------- Sheet2 行级汇总 ----------
    ws2 = wb.create_sheet("行级汇总")
    h2 = ["row_id"] + src_cols + ["gate_lang", "gate_charset", "gate_reason",
                                   "n_sents", "中位句长", "max句长", "句式中疑问", "句式中感叹", "NSW句数"]
    _header(ws2, h2)
    gl_col2 = h2.index("gate_lang") + 1
    gc_col2 = h2.index("gate_charset") + 1
    rr = 1
    by_row: dict = {}
    for s in sentences:
        by_row.setdefault(s.get("row_id"), []).append(s)
    for r_ in records:
        sid = r_["row_id"]
        sents = by_row.get(sid, [])
        gate = r_.get("gate", {}) or {}
        gl = gate.get("lang") or {}
        gc = gate.get("charset") or {}
        row = [sid]
        fields = r_.get("fields") or []
        for idx, c in enumerate(src_cols):
            v = fields[idx] if idx < len(fields) else ""
            if c in numeric_cols:
                v = _num(v)
            row.append(v)
        lens = sorted(s.get("n_chars", 0) for s in sents)
        # 标准中位数（偶数取中间两数均值），与 stats.quantiles 的 p50 同口径；
        # 原先 len(lens)//2 是上中位，与「维度分布」sheet 的 p50 会对不上。
        med = round(statistics.median(lens), 1) if lens else 0
        row += [
            gl.get("status", ""), gc.get("status", ""),
            "; ".join(x for x in [gl.get("reason"), gc.get("reason")] if x),
            len(sents), med, max(lens) if lens else 0,
            sum(1 for s in sents if s.get("sentence_type") == "疑问"),
            sum(1 for s in sents if s.get("sentence_type") == "感叹"),
            sum(1 for s in sents if s.get("nsw_types")),
        ]
        ws2.append(row)
        rr += 1
        _tag(ws2, rr, gl_col2, gl.get("status", ""))
        _tag(ws2, rr, gc_col2, gc.get("status", ""))
    _col_widths(ws2, {"A": 10})

    # ---------- Sheet3 维度分布 ----------
    ws3 = wb.create_sheet("维度分布")

    def block(title, rows):
        ws3.append([title])
        ws3.cell(ws3.max_row, 1).font = Font(bold=True, size=12)
        for r_ in rows:
            ws3.append(r_)

    gate = dims.get("gate", {}) or {}
    length = dims.get("length", {}) or {}
    gate_rows = [
        ["项", "值"],
        ["行数", gate.get("rows") or len(records)],
        ["语言判定 pass 行数", gate.get("lang_pass")],
        ["语言判定 pass 占比", gate.get("lang_pass_ratio")],
        ["语言判定 warn 行数", gate.get("lang_warn")],
        ["语言判定 fail 行数", gate.get("lang_fail")],
        ["字符检查 pass 行数", gate.get("charset_pass")],
        ["字符检查 pass 占比", gate.get("charset_pass_ratio")],
        ["字符检查 warn 行数", gate.get("charset_warn")],
        ["字符检查 fail 行数", gate.get("charset_fail")],
        ["带 emoji 行占比", gate.get("emoji_ratio")],
    ]
    # 降级必须显式可见：否则非目标语用他语资源跑出的干净数字会被当作本语言结论
    if gate.get("degraded"):
        gate_rows.append(["⚠ 口径降级", "是（" + "；".join(gate.get("degraded_reasons") or []) + "）"])
    block("〇、能力与口径（本次计算建立在什么之上）",
          build_capability_rows(cfg, gate, length, dims.get("phoneme", {}) or {}, run_params))
    ws3.append([])
    block("一、合法性 gate", gate_rows)
    ws3.append([])

    # 越界/噪声字符 Top（原实现只存在 records.jsonl，报告层从不展示；此处聚合到维度分布）
    _fc, _nc, _n_frow, _n_nrow = Counter(), Counter(), 0, 0
    for r_ in records:
        gc = (r_.get("gate") or {}).get("charset") or {}
        fch = gc.get("foreign_chars") or []
        nch = gc.get("noise_chars") or []
        if fch:
            _n_frow += 1
        if nch:
            _n_nrow += 1
        for c, n in fch:
            _fc[c] += n
        for c, n in nch:
            _nc[c] += n
    if _fc or _nc:
        block("一之二、越界/噪声字符（全批聚合 Top20）", [
            ["项", "值", "涉及行数"],
            ["越界字符 Top", "、".join(f"{c}:{n}" for c, n in _fc.most_common(20)), _n_frow],
            ["噪声字符 Top", "、".join(f"{c!r}:{n}" for c, n in _nc.most_common(20)), _n_nrow],
            ["越界字符种类", len(_fc), ""],
            ["噪声字符种类", len(_nc), ""],
        ])
        ws3.append([])

    length_syll_na = (length.get("syll") is None)
    syll_cell = ("N/A（" + str(length.get("syll_note") or "口径未定稿") + "）" if length_syll_na
                 else f"{length.get('syll', {}).get('p50')} / {length.get('syll', {}).get('p95')} / {length.get('syll', {}).get('max')}")
    block("二、长度分布（句级）", [
        ["指标", "值"],
        ["句子总数", length.get("n_sents")],
        ["字符口径", "字素簇" if length.get("char_unit") == "grapheme" else "码点 len()"],
        ["单句字符 中位/P95/max", f"{length.get('char', {}).get('p50')} / {length.get('char', {}).get('p95')} / {length.get('char', {}).get('max')}"],
        ["单句词数 中位/P95/max", f"{length.get('word', {}).get('p50')} / {length.get('word', {}).get('p95')} / {length.get('word', {}).get('max')}"],
        ["单句音节 中位/P95/max", syll_cell],
        ["音节口径", str(length.get("syll_method") or "") +
                     (f"（命中率 {length['syll_coverage'].get('dict_hit_ratio')}）"
                      if (length.get("syll_coverage") or {}).get("dict_hit_ratio") is not None else "")],
        ["超 300 字符句数", length.get("over_300")],
        ["超 30 词句数", length.get("over_30w")],
        ["超 60 音节句数", length.get("over_60syll")],
    ])
    ws3.append([])

    st = dims.get("sentence", {}) or {}
    sents_n = st.get("n", 0) or (length.get("n_sents") or 0)
    block("三、句式分布（含目标偏差告警）", [
        ["句式", "句数", "占比", "目标区间", "差距", "告警"],
    ] + [
        [d.get("句式"), d.get("句数"), d.get("占比"),
         (f"{d['目标下']:.0%}–{d['目标上']:.0%}" if "目标下" in d else ""),
         d.get("差距", ""), d.get("告警", "")]
        for d in (st.get("distribution") or [
            {"句式": k, "句数": v, "占比": (round(v / sents_n, 4) if sents_n else 0.0)}
            for k, v in (st.get("types") or {}).items()
        ])
    ])
    ws3.append(["整体判定", st.get("status", ""), f"告警 {st.get('n_warn', 0)} 项"])
    ws3.append([])

    voc = dims.get("vocab", {}) or {}
    block("四、词汇分布", [
        ["指标", "值"],
        ["词例 token 总数", voc.get("tokens")],
        ["型符 unique", voc.get("types")],
        ["型符/词例", voc.get("ttr")],
        ["H_norm", voc.get("h_norm")],
        ["Gini", voc.get("gini")],
        ["Top10% 型符 token 份额", voc.get("top10_share")],
        ["ipa-dict 命中 token 占比", voc.get("dict_hit_ratio")],
        ["虚词 token 占比", voc.get("fw_ratio") if voc.get("fw_ratio") is not None else voc.get("sw_ratio")],
        ["英语功能词 token 占比（近似）", voc.get("en_ratio")],
        ["NSW 命中句占比", voc.get("nsw_sent_ratio")],
        ["高频词 Top20", "、".join(f"{w}:{c}" for w, c in (voc.get("top", []) or [])[:20])],
    ])
    if voc.get("top_truncated"):
        ws3.append(["词频表截断", f"仅列 Top{voc.get('top_n')}（型符共 {voc.get('types')}）；"
                                  f"全量词频见 {voc.get('top_full_source')}"])
    ws3.append([])

    # NSW 分类型明细（V3.4：类别 → 命中句数 → 命中串次数 → 示例）
    if voc.get("nsw_by_category"):
        block("四之二、NSW 分类型（命中具体串）", [
            ["类别", "命中句数", "命中串次数", "示例（串:频次，Top）"],
        ] + [
            [c.get("category"), c.get("sentences"), c.get("hits"),
             "、".join(f"{w}:{n}" for w, n in (c.get("top") or [])[:12])]
            for c in voc["nsw_by_category"]
        ])
        ws3.append([])

    ph = dims.get("phoneme", {}) or {}
    ptr = ph.get("phonotactics") or {}
    pau = ph.get("g2p_audit") or {}
    block("五、音素分布", [
        ["指标", "值"],
        ["词级单音素集合数", ph.get("n_single")],
        ["双音素组合数", ph.get("n_bi")],
        ["双音素总出现", ph.get("total_bi")],
        ["长尾双音素(<3次)数", ph.get("n_longtail_bi")],
        ["字母总数", (ph.get("letters") or {}).get("total")],
        ["字母种类", (ph.get("letters") or {}).get("n_types")],
        ["G2P 未覆盖 token 数", ph.get("g2p_miss")],
        ["　└ 非词形（数字/缩写/符号，本就不音素化）", (ph.get("g2p_miss_breakdown") or {}).get("not_a_word_shape")],
        ["　└ 无兜底器（g2p.fallback=none）", (ph.get("g2p_miss_breakdown") or {}).get("no_fallback_configured")],
        ["　└ epitran 故障", (ph.get("g2p_miss_breakdown") or {}).get("epitran_failed")],
        ["　└ epitran 故障信息", (ph.get("g2p_miss_breakdown") or {}).get("epitran_error")],
        ["词级 IPA 覆盖", f"{sum(1 for r_ in word_rows if r_['ipa'])} / {len(word_rows)}"],
        ["单音素（归一后，异写归并+前鼻化拆分）", ph.get("n_single_norm")],
        ["双音素（归一后）", ph.get("n_bi_norm")],
    ])
    ws3.append([])

    block("五之二、音素组合空间 / 缺失组合（V4.3）", [
        ["指标", "值"],
        ["理论组合空间（native）", ptr.get("n_space")],
        ["语料观测组合", ptr.get("n_observed")],
        ["缺失组合数（空间内未覆盖）", ptr.get("n_missing")],
        ["越界组合数（外来丛，不计缺失）", ptr.get("n_out_of_space")],
        ["缺失组合清单（选句补齐候选）", "、".join(m["combo"] for m in (ptr.get("missing") or [])[:60])],
        ["越界组合 Top（多为英语/阿拉伯借词丛）", "、".join(f"{m['combo']}:{m['count']}" for m in (ptr.get("out_of_space") or [])[:20])],
        ["越界组合总数 / 总出现", f"{ptr.get('n_out_of_space')} / {ptr.get('out_of_space_total')}"
                                  f"（此处仅列 Top20，全量见 dims/phoneme.json）"],
    ])
    ws3.append([])

    block("五之三、G2P 一致性抽验（V4.4）", [
        ["指标", "值"],
        ["抽样词数（词典命中词）", pau.get("n_sample")],
        ["完全一致率", pau.get("rate_exact")],
        ["音素序列一致率（归一后）", pau.get("rate_phoneme")],
        ["折合表示层差异后一致率", pau.get("rate_folded")],
        ["阈值", pau.get("threshold")],
        ["判定", pau.get("status")],
        ["仍不一致（真差异）数", pau.get("n_mismatch")],
    ])
    ws3.append([])

    ety = dims.get("etymology", {}) or {}
    block("六、词源分布（LLM，词级去重）", [
        ["指标", "值"],
        ["去重后词数", ety.get("unique_words")],
        ["已打标词数", ety.get("tagged")],
        ["待打标词数", ety.get("untagged")],
        ["打标失败词数", ety.get("failed")],
        ["残留词数（非本批语料，不计入统计）", ety.get("stale")],
        ["LLM 调用批次数", ety.get("calls")],
        ["批大小", ety.get("batch_size")],
    ])
    ws3.append(["origin", "token 加权", "型符"])
    for k in sorted(set(ety.get("origin") or {}) | set(ety.get("origin_words") or {})):
        ws3.append([k, (ety.get("origin") or {}).get(k, 0), (ety.get("origin_words") or {}).get(k, 0)])

    # ---------- Sheet4 字段说明 ----------
    ws4 = wb.create_sheet("字段说明")
    _header(ws4, ["类别", "字段名", "说明"])
    overrides = _desc_overrides(cfg)
    for c in src_cols:
        ws4.append(["原表字段", c, ORIGINAL_COLUMNS.get(c, "原回流表字段（透传）")])
    for name, desc, dtype in NEW_COLUMNS:
        ws4.append(["新增字段", name, f"{overrides.get(name, desc)}（{dtype}）"])
    for name, desc, dtype in WORD_COLUMNS:
        ws4.append(["词级明细字段", name, f"{overrides.get(name, desc)}（{dtype}）"])
    for nm, desc in [
        ("音素组合", "归一后双音素（bigram）计数表；归一=异写归并(g/ɡ、r/ɾ)+前鼻化拆分"),
        ("缺失组合", "理论音素组合空间内、语料未覆盖的组合（选句补齐候选）"),
        ("G2P抽验", "G2P 一致性抽验不一致清单（词/词典IPA/epitranIPA/差异类型）"),
    ]:
        ws4.append(["附加 sheet", nm, desc])

    # ---------- Sheet5 词级明细 ----------
    ws5 = wb.create_sheet("词级明细")
    _header(ws5, [c[0] for c in WORD_COLUMNS])
    for row in word_rows:
        ws5.append([row["word"], row["ipa"], row["phonemes"], row["n_phonemes"], row["origin"], row["count"]])
    _col_widths(ws5, {"A": 18, "C": 30})
    ws5.freeze_panes = "A2"

    # ---------- Sheet6 音素组合 ----------
    ws6 = wb.create_sheet("音素组合")
    _header(ws6, ["组合", "次数"])
    for m in (ptr.get("observed") or []):
        ws6.append([m["combo"], m["count"]])
    _col_widths(ws6, {"A": 16})

    # ---------- Sheet7 缺失组合 ----------
    ws7 = wb.create_sheet("缺失组合")
    ws7.append(["理论空间大小", ptr.get("n_space")])
    ws7.append(["语料观测组合", ptr.get("n_observed")])
    ws7.append(["缺失组合数", ptr.get("n_missing")])
    ws7.append(["越界组合数（外来丛）", ptr.get("n_out_of_space")])
    ws7.append([])
    ws7.append(["缺失组合（选句补齐候选）", "由", "至"])
    # 样式取 append 之后的 max_row：空行（append([])）不产生单元格、不计入 max_row，
    # 故不能用「append 前 max_row + 1」推算表头行。
    for c in ws7[ws7.max_row]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
    for m in (ptr.get("missing") or []):
        ws7.append([m["combo"], m["from"], m["to"]])
    _col_widths(ws7, {"A": 22})

    # ---------- Sheet8 G2P 抽验 ----------
    ws8 = wb.create_sheet("G2P抽验")
    _header(ws8, ["word", "dict_ipa", "epitran_ipa", "kind"])
    for m in (pau.get("mismatches") or []):
        ws8.append([m.get("word"), m.get("dict_ipa"), m.get("epitran_ipa"), m.get("kind")])
    _col_widths(ws8, {"A": 20, "B": 24, "C": 24})

    # ---------- 落盘 CSV（句子级明细） ----------
    csv_path = outdir / "数据集_全字段.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        # 从第 2 行起：ws 第 1 行就是表头，再迭代会把它当数据写第二遍
        for row in ws.iter_rows(min_row=2, values_only=True):
            w.writerow(row)

    # ---------- 落盘 CSV + JSONL（词级明细） ----------
    word_csv_path = outdir / "数据集_词级明细.csv"
    with word_csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow([c[0] for c in WORD_COLUMNS])
        for row in word_rows:
            w.writerow([row["word"], row["ipa"], row["phonemes"], row["n_phonemes"], row["origin"], row["count"]])
    word_jsonl_path = outdir / "词级明细.jsonl"
    write_jsonl(word_jsonl_path, word_rows)

    # ---------- manifest ----------
    manifest = {
        "run_id": run_params.get("run_id"),
        "lang": run_params.get("lang"),
        "script_family": (getattr(cfg, "script", None) if cfg else None),
        "seed": 0,
        "version": {"corpus_kit": __version__, "python": sys.version.split()[0]},
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "input": {
            "file": str(run_params.get("input")),
            "sha256": _hash_file(Path(run_params["input"])) if run_params.get("input") and Path(run_params["input"]).exists() else None,
            "meta_columns": src_cols,
            "rows": len(records),
            "sentences": len(sentences),
            "unique_words": len(word_rows),
            "words_with_ipa": sum(1 for r_ in word_rows if r_["ipa"]),
            "words_with_origin": sum(1 for r_ in word_rows if r_["origin"]),
            "missing_combos": (ph.get("phonotactics") or {}).get("n_missing"),
            "g2p_audit_status": (ph.get("g2p_audit") or {}).get("status"),
            "sentence_status": (dims.get("sentence") or {}).get("status"),
            "degraded": bool(gate.get("degraded")),
            "degraded_reasons": list(gate.get("degraded_reasons") or []),
            "capability_gaps": list(gate.get("capability_gaps") or []),
            "lang_check": gate.get("lang_check"),
        },
        "params": {k: v for k, v in run_params.items() if k not in ("input",)},
    }
    (outdir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    xlsx_path = outdir / "数据集_全字段.xlsx"
    wb.save(xlsx_path)
    return {"xlsx": str(xlsx_path), "csv": str(csv_path),
            "word_csv": str(word_csv_path), "word_jsonl": str(word_jsonl_path),
            "manifest": str(outdir / "manifest.json")}


def main(argv=None):
    ap = argparse.ArgumentParser(description="report —— 汇总中间态出全字段报告")
    ap.add_argument("--out", required=True, help="output 运行目录")
    ap.add_argument("--meta", default=None, help="INPUT.meta.json 路径")
    ap.add_argument("--run-id", default=None)
    ap.add_argument("--lang", required=True, help="language code or auto")
    ap.add_argument("--input", default="")
    ap.add_argument("--params", default=None,
                    help="本次运行的实际参数（JSON 字符串，如 run_pipeline 透传）；记入 manifest.params")
    args = ap.parse_args(argv)
    outdir = ensure_dir(args.out)
    meta = {}
    meta_src = args.meta if (args.meta and Path(args.meta).exists()) else (outdir / "input_meta.json")
    if Path(meta_src).exists():
        meta = json.loads(Path(meta_src).read_text(encoding="utf-8-sig"))
    sentences = read_jsonl(outdir / "sentences.jsonl", name="sentences.jsonl")
    records = read_jsonl(outdir / "records.jsonl", name="records.jsonl")
    dims = {}
    for name in ("gate", "length", "sentence", "vocab", "phoneme", "etymology"):
        p = outdir / "dims" / f"{name}.json"
        if p.exists():
            dims[name] = json.loads(p.read_text(encoding="utf-8-sig"))
    run_params = {
        "run_id": args.run_id or datetime.now().strftime("%Y%m%d_%H%M%S"),
        "lang": args.lang, "input": args.input,
    }
    if args.params:
        try:
            extra = json.loads(args.params)
            if isinstance(extra, dict):
                run_params["run_args"] = extra
        except json.JSONDecodeError:
            print(f"[report] 警告：--params 不是合法 JSON，已忽略：{args.params[:200]}", file=sys.stderr)
    cfg = None
    if args.lang and args.lang != "auto":
        try:
            from .config import Config
            cfg = Config(args.lang)
        except Exception:
            cfg = None
    elif args.lang == "auto":
        # gate 已把实际选中的语言写进 dims/gate.json，用回读值构造 cfg
        try:
            from .config import Config
            g = dims.get("gate") or {}
            if g.get("lang"):
                cfg = Config(g["lang"])
                run_params["lang"] = g["lang"]
        except Exception:
            cfg = None
    out = build_report(outdir, meta, sentences, records, dims, run_params, cfg=cfg)
    print(json.dumps(out, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())